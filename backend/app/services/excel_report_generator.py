import io
import json
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelReportGenerator:
    """
    Production-Grade Operational Mission Excel Detection Data Report Generator.
    Produces a single, clean, comprehensively styled master data sheet containing
    every acoustic detection record for the selected mission.
    """

    # Maritime / Defense Color Palette
    HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")       # Dark Navy Slate
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")        # Subtle Slate Gray
    CONFIRMED_FILL = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")    # Soft Green
    PENDING_FILL = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")      # Soft Amber
    REJECTED_FILL = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")     # Soft Red

    # Fonts
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")
    REGULAR_FONT = Font(name="Calibri", size=10, color="1E293B")
    MONO_FONT = Font(name="Consolas", size=9, color="334155")
    STATUS_CONFIRMED_FONT = Font(name="Calibri", size=10, bold=True, color="065F46")
    STATUS_PENDING_FONT = Font(name="Calibri", size=10, bold=True, color="92400E")
    STATUS_REJECTED_FONT = Font(name="Calibri", size=10, bold=True, color="991B1B")

    # Borders
    THIN_BORDER = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    HEADER_BORDER = Border(
        left=Side(style='thin', color='1E293B'),
        right=Side(style='thin', color='1E293B'),
        top=Side(style='medium', color='0F172A'),
        bottom=Side(style='medium', color='0B132B')
    )

    LABEL_MAP = {
        "debris_net": "Derelict Ghost Net",
        "pipe_cylinder": "Pipeline / Cylinder",
        "wreck_structure": "Shipwreck Structure",
        "cargo_container": "Cargo Container",
        "naval_mine": "Acoustic Mine Anomaly",
        "pipe_joint": "Pipeline Joint / Free-Span",
        "metal_debris": "Metallic Debris",
        "tire_debris": "Submerged Tire / Rubber",
        "tire": "Submerged Tire / Rubber"
    }

    @classmethod
    def get_display_label(cls, class_name: Optional[str]) -> str:
        if not class_name:
            return "Unclassified Anomaly"
        key = str(class_name).lower().strip()
        return cls.LABEL_MAP.get(key, key.replace("_", " ").title())

    @classmethod
    def _format_timestamp(cls, ts_raw: Any) -> str:
        if not ts_raw:
            return "N/A"
        if isinstance(ts_raw, datetime):
            return ts_raw.strftime("%Y-%m-%d %H:%M:%S UTC")
        ts_str = str(ts_raw).replace("Z", "+00:00")
        try:
            dt = datetime.fromisoformat(ts_str)
            return dt.strftime("%Y-%m-%d %H:%M:%S UTC")
        except Exception:
            return str(ts_raw)

    @classmethod
    def generate_excel(
        cls,
        analysis: Dict[str, Any],
        detections: Optional[List[Dict[str, Any]]] = None
    ) -> bytes:
        """
        Generates a clean, detailed, single-sheet Excel data report containing ALL
        detection records for the selected survey mission.
        """
        wb = openpyxl.Workbook()
        # Remove default sheet
        if wb.active:
            wb.remove(wb.active)

        mission = analysis.get("mission", {})
        m_id = mission.get("mission_id", "MISSION-001")

        # -------------------------------------------------------------------------
        # 1. Resolve Detection Records (Include ALL Mission Targets & Observations)
        # -------------------------------------------------------------------------
        targets = analysis.get("targets") or analysis.get("priority_targets") or []
        if not targets and m_id:
            try:
                from app.db.repository import repo
                targets = repo.get_consolidated_targets(mission_id=m_id)
            except Exception:
                pass

        candidate_dets = detections if detections is not None else analysis.get("detections")
        if candidate_dets is None and m_id:
            try:
                from app.db.repository import repo
                candidate_dets = repo.get_all_detections(mission_id=m_id)
            except Exception:
                pass

        det_list: List[Dict[str, Any]] = []
        covered_target_ids = set()

        # Step 1: Add all observations from detections list
        if candidate_dets:
            for d in candidate_dets:
                det_list.append(d)
                tid = d.get("target_id")
                if tid:
                    covered_target_ids.add(tid)

        # Step 2: Ensure EVERY target in the mission is included!
        for t in (targets or []):
            tid = t.get("target_id") or t.get("id")
            if tid not in covered_target_ids:
                obs_list = t.get("observations") or []
                if obs_list:
                    for o in obs_list:
                        det_list.append(o)
                else:
                    det_list.append(t)
                if tid:
                    covered_target_ids.add(tid)

        # Sort detections by pass number, then target_id, then timestamp if available
        def sort_key(d: Dict[str, Any]):
            return (
                d.get("pass_number") or 1,
                str(d.get("target_id") or ""),
                str(d.get("timestamp") or "")
            )
        try:
            det_list.sort(key=sort_key)
        except Exception:
            pass

        # -------------------------------------------------------------------------
        # 2. Setup Single Main Worksheet: "Mission Detections"
        # -------------------------------------------------------------------------
        ws = wb.create_sheet(title="Mission Detections")
        ws.sheet_properties.tabColor = "0F172A"
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        # Define 21 comprehensive column headers
        headers = [
            ("#", 6),
            ("Mission ID", 16),
            ("Target ID", 16),
            ("Detection ID", 16),
            ("Class", 16),
            ("Classification Label", 24),
            ("Confidence (%)", 16),
            ("Latitude (WGS84)", 18),
            ("Longitude (WGS84)", 18),
            ("Estimated Size (m)", 18),
            ("Shadow Verified", 16),
            ("Review Status", 16),
            ("Timestamp (UTC)", 22),
            ("Pass Number", 14),
            ("Survey Leg", 22),
            ("Sonar Image Reference", 36),
            ("Image ID", 24),
            ("Bounding Box", 34),
            ("Segmentation Data", 38),
            ("Mask Reference", 18),
            ("User ID", 16),
        ]

        header_titles = [h[0] for h in headers]
        default_widths = {idx + 1: h[1] for idx, h in enumerate(headers)}

        ws.append(header_titles)
        ws.row_dimensions[1].height = 28

        # Style Header Row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.border = cls.HEADER_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # -------------------------------------------------------------------------
        # 3. Populate Detection Data Rows
        # -------------------------------------------------------------------------
        row_idx = 2
        for idx, det in enumerate(det_list, start=1):
            # Resolve fields
            mission_val = det.get("mission_id") or m_id or "N/A"
            target_val = det.get("target_id") or "N/A"
            det_id_val = det.get("id") or det.get("observation_id") or f"det_{idx}"
            cls_name = det.get("class") or det.get("category") or det.get("target_class") or "unclassified"
            label_val = det.get("label") or cls.get_display_label(cls_name)

            # Confidence as numeric float (0.0 to 1.0)
            raw_conf = det.get("confidence") if det.get("confidence") is not None else det.get("fused_confidence")
            conf_val: Optional[float] = None
            if raw_conf is not None:
                try:
                    c_float = float(raw_conf)
                    conf_val = c_float / 100.0 if c_float > 1.0 else c_float
                except (ValueError, TypeError):
                    conf_val = None

            # Coordinates
            lat_val: Optional[float] = None
            if det.get("latitude") is not None:
                try:
                    lat_val = float(det["latitude"])
                except (ValueError, TypeError):
                    pass

            lon_val: Optional[float] = None
            if det.get("longitude") is not None:
                try:
                    lon_val = float(det["longitude"])
                except (ValueError, TypeError):
                    pass

            # Estimated size (defaults to 3.0m if 0 or missing)
            raw_size = det.get("estimated_size_m")
            size_val: float = 3.0
            if raw_size is not None:
                try:
                    s_float = float(raw_size)
                    size_val = s_float if s_float > 0 else 3.0
                except (ValueError, TypeError):
                    size_val = 3.0

            # Shadow verification
            sv = det.get("shadow_verified")
            if sv is True or sv == 1 or str(sv).lower() == "true":
                shadow_str = "VERIFIED"
            elif sv is False or sv == 0 or str(sv).lower() == "false":
                shadow_str = "UNVERIFIED"
            else:
                shadow_str = "UNKNOWN"

            # Status
            raw_stat = det.get("status") or det.get("human_review_status") or "pending_review"
            norm_stat = str(raw_stat).lower().strip()
            if norm_stat in ["confirmed", "verified"]:
                status_display = "Confirmed"
            elif norm_stat in ["rejected", "false_alarm"]:
                status_display = "Rejected"
            else:
                status_display = "Pending Review"

            # Timestamp
            ts_display = cls._format_timestamp(det.get("timestamp") or det.get("created_at"))

            # Pass Number and Survey Leg
            pass_num = det.get("pass_number", 1)
            survey_leg = det.get("survey_leg") or f"Swath Pass {pass_num}"

            # Images
            sonar_img = det.get("sonar_image_ref") or "None"
            image_id = det.get("image_id") or "None"

            # Bounding Box JSON
            bbox_raw = det.get("bounding_box")
            if bbox_raw is not None:
                bbox_str = json.dumps(bbox_raw) if not isinstance(bbox_raw, str) else bbox_raw
            else:
                bbox_str = "N/A"

            # Segmentation Data JSON
            seg_raw = det.get("segmentation")
            if seg_raw is not None:
                seg_str = json.dumps(seg_raw) if not isinstance(seg_raw, str) else seg_raw
            else:
                seg_str = "N/A"

            # Mask Reference
            mask_ref = det.get("mask_ref") or "None"

            # User ID
            user_id = det.get("user_id") or "N/A"

            row_data = [
                idx,
                mission_val,
                target_val,
                det_id_val,
                cls_name,
                label_val,
                conf_val,
                lat_val,
                lon_val,
                size_val,
                shadow_str,
                status_display,
                ts_display,
                pass_num,
                survey_leg,
                sonar_img,
                image_id,
                bbox_str,
                seg_str,
                mask_ref,
                user_id
            ]

            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20

            # Apply row styling
            is_even = (row_idx % 2 == 0)
            row_fill = cls.ZEBRA_FILL if is_even else PatternFill(fill_type=None)

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = cls.THIN_BORDER
                c.font = cls.REGULAR_FONT

                if row_fill.fill_type:
                    c.fill = row_fill

                # Specific Column Alignments and Formats
                if col_idx == 1:  # #
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in [2, 3, 4, 14, 17, 20, 21]:  # IDs, Pass Number
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 7:  # Confidence (%)
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    if conf_val is not None:
                        c.number_format = "0.0%"
                elif col_idx in [8, 9]:  # Coordinates
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    if c.value is not None and isinstance(c.value, (int, float)):
                        c.number_format = "0.000000"
                elif col_idx == 10:  # Estimated Size
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    if size_val is not None:
                        c.number_format = "0.00"
                elif col_idx == 11:  # Shadow Verified
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    if shadow_str == "VERIFIED":
                        c.font = cls.STATUS_CONFIRMED_FONT
                    elif shadow_str == "UNVERIFIED":
                        c.font = cls.STATUS_REJECTED_FONT
                elif col_idx == 12:  # Review Status
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    if status_display == "Confirmed":
                        c.font = cls.STATUS_CONFIRMED_FONT
                        c.fill = cls.CONFIRMED_FILL
                    elif status_display == "Rejected":
                        c.font = cls.STATUS_REJECTED_FONT
                        c.fill = cls.REJECTED_FILL
                    else:
                        c.font = cls.STATUS_PENDING_FONT
                        c.fill = cls.PENDING_FILL
                elif col_idx == 13:  # Timestamp
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 16:  # Sonar Image Ref
                    c.alignment = Alignment(horizontal="left", vertical="center")
                elif col_idx in [18, 19]:  # Bounding Box & Segmentation JSON
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    c.font = cls.MONO_FONT
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

            row_idx += 1

        # -------------------------------------------------------------------------
        # 4. Handle Empty Detection Fallback
        # -------------------------------------------------------------------------
        if len(det_list) == 0:
            ws.append([1, m_id, "NO_RECORDS", "N/A", "N/A", "No Detections Found", None, None, None, None, "UNKNOWN", "N/A", "N/A", 1, "N/A", "None", "None", "N/A", "N/A", "None", "N/A"])
            ws.row_dimensions[2].height = 22
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=2, column=col_idx)
                c.border = cls.THIN_BORDER
                c.font = cls.REGULAR_FONT
                c.alignment = Alignment(horizontal="center", vertical="center")
            row_idx = 3

        # -------------------------------------------------------------------------
        # 5. Column Widths & AutoFilter
        # -------------------------------------------------------------------------
        max_row = max(2, row_idx - 1)
        max_col_letter = get_column_letter(len(headers))

        # Enable AutoFilter on the table
        ws.auto_filter.ref = f"A1:{max_col_letter}{max_row}"

        # Set column widths with smart clamping
        for col_idx, (hdr_name, def_w) in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(hdr_name)

            # Sample first 50 rows to calculate auto width without freezing on huge segmentation polygons
            for r in range(1, min(max_row + 1, 50)):
                val_str = str(ws.cell(row=r, column=col_idx).value or "")
                # Cap line length if contains newline or JSON
                first_line = val_str.split("\n")[0]
                max_len = max(max_len, min(len(first_line), 45))

            # Constrain JSON & URL columns to designated width
            if col_idx in [16, 18, 19]:
                ws.column_dimensions[col_letter].width = def_w
            else:
                ws.column_dimensions[col_letter].width = max(def_w, max_len + 3)

        # -------------------------------------------------------------------------
        # 6. Save and Return In-Memory Bytes
        # -------------------------------------------------------------------------
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


excel_report_generator = ExcelReportGenerator()
